from pathlib import Path

from docx import Document
from docx.enum.text import WD_COLOR_INDEX


ROOT = Path(r"d:\User\hqy\Downloads\ensptomcpfortrae-main\ensptomcpfortrae-main")
SRC = ROOT / "高校无线网络方案设计与工程实践_文字符号修订版_当前eNSP一致版.docx"
OUT = ROOT / "高校无线网络方案设计与工程实践_文字符号修订版_当前eNSP一致版_高亮标注版.docx"


PARAGRAPH_MATCHES = [
    "FIT-AP+AC 三层组网与本地转发配置",
    "VLAN 规划、管理网络连通性、DHCP 地址分配和 SSID 配置逻辑能够得到验证",
    "eNSP 平台本身不提供原生 OLT、OBD 和 ONU 设备模型",
    "仿真环境未直接复现真实 PON 协议过程",
    "仿真验证重点定位为“管理网络、地址规划、业务 VLAN、AC 管理与终端接入逻辑验证”",
    "针对 XGPON 接入部分，本文在方案设计层面引入其承载思路",
    "针对 XGPON 接入部分，由于平台不支持原生 OLT、OBD 和 ONU 设备模型",
]


def highlight_paragraph(paragraph) -> bool:
    if not paragraph.text.strip():
        return False
    changed = False
    for run in paragraph.runs:
        run.font.highlight_color = WD_COLOR_INDEX.YELLOW
        changed = True
    return changed


def highlight_cell(cell) -> bool:
    changed = False
    for paragraph in cell.paragraphs:
        changed = highlight_paragraph(paragraph) or changed
    return changed


def main() -> None:
    doc = Document(str(SRC))
    highlighted = 0

    for paragraph in doc.paragraphs:
        if any(marker in paragraph.text for marker in PARAGRAPH_MATCHES):
            if highlight_paragraph(paragraph):
                highlighted += 1

    for table in doc.tables:
        header = table.cell(0, 0).text.strip() if table.rows and table.columns else ""

        if header == "对象":
            for row in table.rows[1:]:
                key = row.cells[0].text.strip()
                if key in {"Core1-HNCU", "Core2-HNCU", "AGG-Dorm-Male"} or "ONU/" in key:
                    for cell in row.cells:
                        if highlight_cell(cell):
                            highlighted += 1

        if header == "SSID":
            for row in table.rows[1:]:
                for cell in row.cells:
                    if highlight_cell(cell):
                        highlighted += 1

        if header == "区域":
            for row in table.rows[1:]:
                for cell in row.cells:
                    if "HNCU-" in cell.text or "规划项" in cell.text:
                        if highlight_cell(cell):
                            highlighted += 1

    doc.save(str(OUT))
    print(OUT.name)
    print(f"highlighted={highlighted}")


if __name__ == "__main__":
    main()
