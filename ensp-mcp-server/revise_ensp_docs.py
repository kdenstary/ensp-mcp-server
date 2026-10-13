from pathlib import Path

from docx import Document
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.util import Inches, Pt


ROOT = Path(r"d:\User\hqy\Downloads\ensptomcpfortrae-main\ensptomcpfortrae-main")
DOC_SRC = ROOT / "高校无线网络方案设计与工程实践_文字符号修订版.docx"
PPT_SRC = ROOT / "高校无线网络方案设计与工程实践_答辩PPT_最终呈现版.pptx"
IMAGE_SRC = ROOT / "图片2.png"
DOC_OUT = ROOT / "高校无线网络方案设计与工程实践_文字符号修订版_当前eNSP一致版.docx"
PPT_OUT = ROOT / "高校无线网络方案设计与工程实践_答辩PPT_最终呈现版_当前eNSP一致版.pptx"


def replace_text(text: str) -> str:
    replacements = {
        "HCU-Teacher": "HNCU-Teacher",
        "HCU-Student": "HNCU-Student",
        "HCU-Guest": "HNCU-Guest",
        "HCU-Dorm": "HNCU-Dorm",
        "HCU-Manage": "HNCU-Manage（规划项）",
        "FIT-AP+AC 三层隧道转发配置": "FIT-AP+AC 三层组网与本地转发配置",
        "结果表明，该方案整体设计思路合理，关键功能能够正常实现，具备较好的工程可行性，可为高校无线网络建设提供一定的参考。": (
            "结果表明，该方案整体设计思路合理，VLAN 规划、管理网络连通性、DHCP 地址分配和 "
            "SSID 配置逻辑能够得到验证；受 eNSP 平台对 AP 节点仿真能力限制，AP 全量自动注册与上线"
            "过程未完全复现，但不影响方案的总体工程可行性。"
        ),
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    return text


def set_font(text_frame, size=12, bold=False, color=RGBColor(31, 41, 55)) -> None:
    for paragraph in text_frame.paragraphs:
        for run in paragraph.runs:
            run.font.size = Pt(size)
            run.font.bold = bold
            run.font.color.rgb = color


def add_textbox(slide, left, top, width, height, text, size=12, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    frame = box.text_frame
    frame.word_wrap = True
    frame.clear()
    frame.text = text
    set_font(frame, size=size, bold=bold)
    return box


def add_labeled_box(slide, left, top, width, height, title, body):
    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor(244, 247, 251)
    shape.line.color.rgb = RGBColor(189, 199, 216)

    frame = shape.text_frame
    frame.clear()
    p1 = frame.paragraphs[0]
    p1.text = title
    p1.runs[0].font.bold = True
    p1.runs[0].font.size = Pt(15)
    p1.runs[0].font.color.rgb = RGBColor(28, 63, 170)

    p2 = frame.add_paragraph()
    p2.text = body
    p2.runs[0].font.size = Pt(12)
    p2.runs[0].font.color.rgb = RGBColor(31, 41, 55)
    return shape


def remove_shapes_by_index(slide, indexes):
    shapes = list(slide.shapes)
    for idx in sorted(indexes, reverse=True):
        shape = shapes[idx]
        shape._element.getparent().remove(shape._element)


def revise_docx() -> None:
    doc = Document(str(DOC_SRC))

    for paragraph in doc.paragraphs:
        new_text = replace_text(paragraph.text)

        if "本方案在楼宇接入侧引入XGPON 承载思路" in new_text:
            new_text = (
                "本方案在设计层面引入 XGPON 楼宇接入承载思路，逻辑链路为“核心交换机—OLT—主干光纤—"
                "OBD 分光器—ONU—PoE 交换机/中心 AP—AP 或 RU—终端用户”。由于 eNSP 平台不支持原生 "
                "OLT、OBD 和 ONU 设备模型，仿真中采用 CE 类交换设备重命名、二层 Trunk 链路和文字标注的"
                "方式进行等效表达，重点验证 VLAN 承载、IP 地址规划、AC 管理、SSID 下发及终端获取地址等关键网络逻辑。"
            )
        elif "其逻辑链路可表示为" in new_text:
            new_text = (
                "需要说明的是，eNSP 平台本身不提供原生 OLT、OBD 和 ONU 设备模型，因此本文未对真实 PON "
                "协议交互过程进行直接仿真。为保证方案表达的完整性与可视化效果，本文在仿真拓扑中采用可用交换设备对象进行重命名，"
                "并通过二层 Trunk 链路和文字标注对 XGPON 接入结构进行等效模拟。该方法主要用于验证宿舍区承载链路中的 VLAN "
                "透传关系、业务接入路径以及与 AC、PoE 交换机、AP/RU 之间的逻辑连接，不影响方案设计本身的工程合理性。"
            )
        elif "在验证环节，基于华为eNSP仿真平台搭建简化测试环境" in new_text:
            new_text = (
                "在验证环节，基于华为 eNSP 仿真平台搭建简化测试环境，对 AC 与 AP 关联、SSID 下发、终端接入、"
                "管理网络连通性和 DHCP 地址分配等关键逻辑进行了验证；针对 XGPON 接入部分，由于平台不支持原生 "
                "OLT、OBD 和 ONU 设备模型，本文采用二层链路与文字标注方式进行等效表达。"
            )
        elif "仿真环境中，核心层采用 Core1-HNCU 与 Core2-HNCU 模拟双核心结构" in new_text:
            new_text = (
                "仿真环境未直接复现真实 PON 协议过程，而是使用 eNSP 可用设备对象对 XGPON 接入链路进行可视化等效模拟。"
                "其中，OLT-Sim 表示汇聚侧光接入汇聚节点，OBD-Splitter 表示无源分光关系，ONU-Dorm 表示楼栋接入侧的业务汇接节点。"
                "该方式主要用于表达承载路径与 VLAN 透传逻辑，不作为真实 PON 协议仿真结果。与此同时，核心层采用 Core1-HNCU "
                "与 Core2-HNCU 模拟双核心结构，校园数据中心部署 Web、FTP 等业务服务器，出口侧通过主备防火墙连接 ISP 网络，"
                "教学、办公和宿舍区域分别通过接入交换机、PoE 交换机和 AC 完成管理与业务链路验证。"
            )
        elif "在拓扑与配置对应关系上，本文保持 VLAN 10 为管理 VLAN" in new_text:
            new_text = (
                "在拓扑与配置对应关系上，本文保持 VLAN 10 为管理 VLAN，VLAN 201、VLAN 202 和 VLAN 203 "
                "分别对应教职工、学生和访客无线业务，VLAN 301、VLAN 401 和 VLAN 501 分别对应教学、宿舍和办公区域业务。"
                "AC 管理地址设置为 10.1.10.100，核心交换机配置对应 Vlanif 网关和 DHCP 地址池。针对宿舍区 XGPON 承载链路，"
                "本文在仿真中采用二层 Trunk 链路、设备重命名和文字标注进行等效模拟，重点验证 VLAN 透传、地址规划和无线管理链路的逻辑可行性。"
            )
        elif "结合前述管理网络连通性测试结果和 DHCP 地址池配置结果分析" in new_text:
            new_text = (
                "结合前述管理网络连通性测试结果和 DHCP 地址池配置结果分析，可以认为本文已在网络侧完成无线接入设备接入管理网所需的主要条件配置，"
                "包括管理 VLAN 划分、三层接口建立、地址池分配和控制器接入验证等。AP 未能进一步完成自动注册，更可能与 eNSP 对 AP 节点的"
                "仿真能力、镜像支持范围及设备模型限制有关，而并不意味着本文提出的总体方案设计存在原则性错误。因此，本文的仿真验证重点定位为"
                "“管理网络、地址规划、业务 VLAN、AC 管理与终端接入逻辑验证”，而非对真实光接入网协议与 AP 全流程注册行为的完整复现。"
            )
        elif new_text.startswith("本毕业设计围绕“湖南城市学院无线网络方案设计与工程实践”这一主题"):
            new_text = (
                "本毕业设计围绕“湖南城市学院无线网络方案设计与工程实践”这一主题，系统完成了需求分析、总体架构设计、详细覆盖规划、"
                "安全认证设计及关键网络逻辑验证。方案采用 Wi-Fi 6、AC+Fit AP 集中管理、敏捷分布式 Wi-Fi 和分层安全体系，"
                "能够较好满足教学区、宿舍区和办公区的差异化接入需求。针对 XGPON 接入部分，本文在方案设计层面引入其承载思路；"
                "在 eNSP 仿真层面，由于平台不支持原生 OLT、OBD 和 ONU 设备模型，因此采用交换设备重命名、二层 Trunk 链路和拓扑标注的"
                "方式进行等效模拟。最终完成了校园无线网络总体方案设计，并在 eNSP 平台上对 VLAN、地址规划、AC 管理、SSID/VAP 映射、"
                "DHCP 地址分配与基础连通性等关键网络逻辑进行了验证，表明方案具有较好的工程参考价值。"
            )

        if new_text != paragraph.text:
            paragraph.text = new_text

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                new_text = replace_text(cell.text)
                if new_text != cell.text:
                    cell.text = new_text

    for table in doc.tables:
        header = table.cell(0, 0).text.strip() if table.rows and table.columns else ""

        if header == "对象":
            for row in table.rows[1:]:
                key = row.cells[0].text.strip()
                if key == "Core1-HNCU":
                    row.cells[2].text = "10.2.1.254/10.2.2.254/10.2.3.254/10.3.1.254/10.4.1.254/10.5.1.254/10.6.0.254"
                    row.cells[3].text = "10/201/202/203/301/401/501/600"
                elif key == "Core2-HNCU":
                    row.cells[2].text = "10.2.1.253/10.2.2.253/10.2.3.253/10.3.1.253/10.4.1.253/10.5.1.253/10.6.0.253"
                    row.cells[3].text = "10/201/202/203/301/401/501/600"
                elif key == "AGG-Dorm-Male":
                    row.cells[3].text = "10/202/203/401"
                elif "ONU/" in key:
                    row.cells[3].text = "10/202/203/401"

        if header == "SSID":
            for row in table.rows[1:]:
                key = row.cells[0].text.strip()
                if key == "HNCU-Manage（规划项）":
                    row.cells[1].text = "运维人员"
                    row.cells[2].text = "账号认证+ACL"
                    row.cells[3].text = "10"
                    row.cells[4].text = "设备管理和运维监控（本次仿真未启用）"

    doc.save(str(DOC_OUT))


def revise_slide_7(slide):
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        if shape.text.strip() == "OLT":
            shape.text = "OLT-Sim"
        elif "OBD" in shape.text:
            shape.text = "OBD-Splitter"
        elif shape.text.strip() == "ONU":
            shape.text = "ONU-Dorm"

    remove_shapes_by_index(slide, list(range(23, 36)))
    add_textbox(slide, Inches(0.7), Inches(0.95), Inches(5.0), Inches(0.35), "XGPON 承载逻辑设计与 eNSP 等效表达", size=13, bold=True)
    add_labeled_box(
        slide,
        Inches(0.75),
        Inches(2.35),
        Inches(4.5),
        Inches(0.95),
        "设计目标",
        "宿舍区采用 XGPON 承载思路，解决楼栋分散、链路距离长的问题。",
    )
    add_labeled_box(
        slide,
        Inches(0.75),
        Inches(3.45),
        Inches(4.5),
        Inches(0.95),
        "平台限制",
        "eNSP 不支持原生 OLT、OBD、ONU 设备模型。",
    )
    add_labeled_box(
        slide,
        Inches(0.75),
        Inches(4.55),
        Inches(4.5),
        Inches(1.05),
        "仿真方式",
        "采用二层链路、设备重命名和文字标注进行等效表达，重点验证 VLAN、地址规划、AC 管理与终端接入逻辑。",
    )
    slide.shapes.add_picture(str(IMAGE_SRC), Inches(5.55), Inches(2.28), width=Inches(5.35))
    add_textbox(slide, Inches(5.65), Inches(5.62), Inches(5.1), Inches(0.3), "图 7-1 eNSP 中宿舍区 XGPON 承载链路等效示意图", size=10)
    add_textbox(
        slide,
        Inches(0.8),
        Inches(6.0),
        Inches(10.2),
        Inches(0.38),
        "说明：eNSP 不支持原生 OLT、OBD、ONU 设备模型，图中相关节点采用交换设备重命名与二层链路方式进行等效模拟，用于表达宿舍区 XGPON 承载逻辑。",
        size=9,
    )


def revise_slide_13(slide):
    for shape in slide.shapes:
        if not hasattr(shape, "text"):
            continue
        if shape.text.strip() == "XGPON/ONU":
            shape.text = "XGPON 等效接入"

    add_textbox(
        slide,
        Inches(0.85),
        Inches(5.08),
        Inches(10.1),
        Inches(0.36),
        "说明：图中 OLT-Sim / OBD-Splitter / ONU-Dorm 为等效标识，不代表 eNSP 原生支持对应设备模型。",
        size=9,
    )


def revise_slide_15(slide):
    remove_shapes_by_index(slide, list(range(6, 32)))
    slide.shapes.add_picture(str(IMAGE_SRC), Inches(0.8), Inches(1.25), width=Inches(7.05))
    add_textbox(slide, Inches(0.95), Inches(4.05), Inches(6.8), Inches(0.3), "图 15-1 eNSP 仿真拓扑与关键验证对象", size=10)
    add_labeled_box(
        slide,
        Inches(8.05),
        Inches(1.32),
        Inches(3.0),
        Inches(1.0),
        "XGPON 承载",
        "采用等效模拟方式表达承载结构，不验证真实 PON 协议过程。",
    )
    add_labeled_box(
        slide,
        Inches(8.05),
        Inches(2.5),
        Inches(3.0),
        Inches(1.15),
        "验证重点",
        "VLAN、DHCP、AC 管理、SSID/VAP 映射与终端接入逻辑。",
    )
    add_textbox(
        slide,
        Inches(8.05),
        Inches(3.82),
        Inches(3.02),
        Inches(0.55),
        "说明：本页验证对象为 VLAN 承载、AC 管理、SSID/VAP 映射、DHCP 地址获取与基础连通性；XGPON 部分仅作逻辑等效表达，不作为真实 PON 协议仿真结果。",
        size=9,
    )


def revise_pptx() -> None:
    prs = Presentation(str(PPT_SRC))

    for slide_index, slide in enumerate(prs.slides, start=1):
        for shape in slide.shapes:
            if not hasattr(shape, "text"):
                continue

            text = shape.text
            new_text = replace_text(text)

            if slide_index == 1 and "汇报人" in text and "指导教师" in text:
                new_text = "汇报人：请替换为本人姓名    指导教师：请替换为教师姓名"

            if slide_index == 9 and "规划思路" in text:
                new_text = (
                    "规划思路：管理IP与业务IP分离，按用户角色和功能区域划分广播域，通过SSID-VLAN映射实现分类接入和权限控制。"
                    "当前 eNSP 核心同时配置 VLAN301、VLAN501；管理SSID为规划项，本次仿真未启用。"
                )

            if slide_index == 15:
                if "SSID下发与终端接入" in text:
                    new_text = "SSID配置与业务映射检查"
                elif "AP正常广播SSID，STA可接入" in text:
                    new_text = "AC已完成SSID/VAP/VLAN配置，AP全量在线受平台限制"
                elif "AP管理状态检查" in text:
                    new_text = "AP注册信息与管理状态检查"
                elif "AC可查看AP在线与管理接口状态" in text:
                    new_text = "AC可查看AP注册信息，当前状态以 idle/fault 为主"

            if new_text != text:
                shape.text = new_text

    through_shapes = [sh for sh in prs.slides[14].shapes if hasattr(sh, "text") and sh.text.strip() == "通过"]
    if len(through_shapes) >= 4:
        through_shapes[2].text = "部分通过"
        through_shapes[3].text = "受平台限制"

    revise_slide_7(prs.slides[6])
    revise_slide_13(prs.slides[12])
    revise_slide_15(prs.slides[14])

    prs.save(str(PPT_OUT))


def main() -> None:
    revise_docx()
    revise_pptx()
    print(DOC_OUT.name)
    print(PPT_OUT.name)


if __name__ == "__main__":
    main()
